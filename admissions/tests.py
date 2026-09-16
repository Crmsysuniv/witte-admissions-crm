import io
import openpyxl
from django.test import TestCase, Client
from django.urls import reverse
from django.contrib.auth import get_user_model
from admissions.models import Faculty, Specialty, EducationProgram, ExamSubject, ExamScore, Application, ApplicationDocument
from accounts.models import ApplicantProfile, OfficerProfile
from admissions.exports import generate_competitive_list_workbook, export_rating_xlsx_response

User = get_user_model()


class ExcelExportCompetitiveListTests(TestCase):
    """
    Тестирование модуля генерации и выгрузки отчетов в формате .xlsx (openpyxl).
    Проверяет:
    - Корректность структуры книги Excel (шапка вуза, реквизиты программы, заголовки);
    - Ранжирование поступающих по сумме баллов и подсчет проходного балла;
    - Выделение зон квоты (зеленая зона плана мест и резерв);
    - Формирование бинарного HTTP-ответа с нужными заголовками Content-Disposition;
    - Доступ к эндпоинтам экспорта со стороны пользователей и приемной комиссии.
    """
    def setUp(self):
        self.client = Client()

        # Сотрудник приемной комиссии
        self.officer_user = User.objects.create_user(
            username='officer_exp',
            password='testpassword123',
            email='officer_exp@witte.ru',
            first_name='Анна',
            last_name='Ильина',
            role=User.Role.OFFICER
        )
        OfficerProfile.objects.create(
            user=self.officer_user,
            position='Специалист по зачислению',
            cabinet='204'
        )

        # Абитуриент 1 (высокий балл + оригинал)
        self.applicant1 = User.objects.create_user(
            username='app_alex',
            password='testpassword123',
            email='alex@example.com',
            first_name='Алексей',
            last_name='Смирнов',
            role=User.Role.APPLICANT
        )
        ApplicantProfile.objects.create(
            user=self.applicant1,
            snils='111-222-333 44'
        )

        # Абитуриент 2 (средний балл + копия)
        self.applicant2 = User.objects.create_user(
            username='app_elena',
            password='testpassword123',
            email='elena@example.com',
            first_name='Елена',
            last_name='Попова',
            role=User.Role.APPLICANT
        )
        ApplicantProfile.objects.create(
            user=self.applicant2,
            snils='555-666-777 88'
        )

        # Факультет и специальность
        self.faculty = Faculty.objects.create(
            name='Факультет информационных технологий',
            code='ФИТ'
        )
        self.specialty = Specialty.objects.create(
            faculty=self.faculty,
            code='09.03.03',
            name='Прикладная информатика',
            education_level=Specialty.EducationLevel.BACHELOR,
            budget_places=1,
            paid_places=10
        )
        self.program = EducationProgram.objects.create(
            specialty=self.specialty,
            study_form=EducationProgram.StudyForm.FULL_TIME,
            tuition_fee=140000.00,
            duration='4 года'
        )

        # Предметы
        self.subj_math = ExamSubject.objects.create(name='Математика', min_score=39)
        self.subj_rus = ExamSubject.objects.create(name='Русский язык', min_score=40)
        self.subj_it = ExamSubject.objects.create(name='Информатика и ИКТ', min_score=44)

        # Заявление 1 (Смирнов: 80 + 90 = 170)
        self.app1 = Application.objects.create(
            applicant=self.applicant1,
            program=self.program,
            status=Application.Status.APPROVED,
            financing_type=Application.FinancingType.BUDGET
        )
        ExamScore.objects.create(application=self.app1, subject=self.subj_rus, score=90, is_verified=True)
        ExamScore.objects.create(application=self.app1, subject=self.subj_math, score=80, is_verified=True)
        
        # Документ оригинал для Смирнова
        ApplicationDocument.objects.create(
            application=self.app1,
            document_type=ApplicationDocument.DocumentType.CERTIFICATE,
            is_verified=True
        )

        # Заявление 2 (Попова: 60 + 70 = 130)
        self.app2 = Application.objects.create(
            applicant=self.applicant2,
            program=self.program,
            status=Application.Status.APPROVED,
            financing_type=Application.FinancingType.BUDGET
        )
        ExamScore.objects.create(application=self.app2, subject=self.subj_rus, score=70, is_verified=True)
        ExamScore.objects.create(application=self.app2, subject=self.subj_math, score=60, is_verified=True)

    def test_generate_competitive_list_workbook_structure(self):
        """Проверка структуры и содержимого сгенерированной книги openpyxl."""
        wb = generate_competitive_list_workbook(
            program=self.program,
            financing_type=Application.FinancingType.BUDGET
        )
        self.assertIn("Конкурсный список", wb.sheetnames)
        ws = wb["Конкурсный список"]

        # Проверка шапки
        self.assertIn("МОСКОВСКИЙ УНИВЕРСИТЕТ ИМЕНИ С.Ю. ВИТТЕ", ws['A1'].value)
        self.assertIn("09.03.03", ws['A3'].value)
        self.assertIn("Прикладная информатика", ws['A3'].value)

        # Проверка заголовков таблицы на строке 7
        self.assertEqual(ws['A7'].value, "№ п/п")
        self.assertEqual(ws['D7'].value, "ФИО абитуриента")
        self.assertEqual(ws['E7'].value, "Сумма баллов")

        # Проверка данных (Смирнов идет первым с суммой 170)
        self.assertEqual(ws['A8'].value, 1)
        self.assertEqual(ws['D8'].value, "Алексей Смирнов")
        self.assertEqual(ws['E8'].value, 170)
        self.assertEqual(ws['I8'].value, "В зоне зачисления")  # Так как budget_places = 1

        # Попова идет второй с суммой 130
        self.assertEqual(ws['A9'].value, 2)
        self.assertEqual(ws['D9'].value, "Елена Попова")
        self.assertEqual(ws['E9'].value, 130)
        self.assertEqual(ws['I9'].value, "В резерве")

    def test_export_rating_xlsx_response_headers(self):
        """Проверка бинарного HTTP-ответа и заголовков выгрузки .xlsx."""
        response = export_rating_xlsx_response(
            program=self.program,
            financing_type=Application.FinancingType.BUDGET
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('.xlsx', response['Content-Disposition'])

        # Проверка валидности zip/xlsx сигнатуры (PK\x03\x04)
        self.assertTrue(response.content.startswith(b'PK'))

        # Проверка чтения книги из байтового потока
        stream = io.BytesIO(response.content)
        loaded_wb = openpyxl.load_workbook(stream)
        self.assertIn("Конкурсный список", loaded_wb.sheetnames)

    def test_export_rating_view_url(self):
        """Проверка эндпоинта /rating/export/ через Django test client."""
        url = reverse('export_rating_root')
        response = self.client.get(url, {
            'program': self.program.id,
            'financing': 'BUDGET'
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

    def test_export_rating_filter_only_originals(self):
        """Проверка выгрузки списка только с оригиналами документов."""
        wb = generate_competitive_list_workbook(
            program=self.program,
            financing_type=Application.FinancingType.BUDGET,
            only_originals=True
        )
        ws = wb["Конкурсный список"]
        # Только Смирнов (у него есть оригинал)
        self.assertEqual(ws['D8'].value, "Алексей Смирнов")
        # На следующей строке уже нет Поповой
        self.assertNotEqual(ws['D9'].value, "Елена Попова")

    def test_generate_application_receipt_docx_content(self):
        """Проверка формирования расписки в формате .docx с помощью python-docx."""
        from admissions.exports import generate_application_receipt_docx
        doc = generate_application_receipt_docx(self.app1)

        # Проверка текста в абзацах
        full_text = "\n".join([p.text for p in doc.paragraphs])
        self.assertIn("МОСКОВСКИЙ УНИВЕРСИТЕТ ИМЕНИ С.Ю. ВИТТЕ", full_text)
        self.assertIn(f"РАСПИСКА В ПРИЕМЕ ДОКУМЕНТОВ № 2026-{self.app1.id:04d}", full_text)

        # Проверка таблиц документа
        tables_text = ""
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    tables_text += " " + cell.text

        self.assertIn("Алексей Смирнов", tables_text)
        self.assertIn("111-222-333 44", tables_text)
        self.assertIn("09.03.03", tables_text)
        self.assertIn("Прикладная информатика", tables_text)
        self.assertIn("Русский язык", tables_text)
        self.assertIn("90 б.", tables_text)
        self.assertIn("170 б.", tables_text)

    def test_export_receipt_docx_response(self):
        """Проверка HTTP-ответа и MIME-типа расписки .docx."""
        from admissions.exports import export_receipt_docx_response
        import docx

        response = export_receipt_docx_response(self.app1)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        self.assertIn('attachment;', response['Content-Disposition'])
        self.assertIn('.docx', response['Content-Disposition'])
        self.assertTrue(response.content.startswith(b'PK'))

        # Проверка считывания документа из байтов
        stream = io.BytesIO(response.content)
        loaded_doc = docx.Document(stream)
        self.assertTrue(len(loaded_doc.paragraphs) > 5)

    def test_student_download_receipt_endpoint(self):
        """Проверка скачивания расписки абитуриентом через личный кабинет."""
        self.client.login(username='app_alex', password='testpassword123')
        url = reverse('student:download_receipt_docx', kwargs={'application_id': self.app1.id})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

    def test_officer_download_receipt_endpoint(self):
        """Проверка формирования расписки сотрудником приемной комиссии."""
        self.client.login(username='officer_exp', password='testpassword123')
        url = reverse('officer:application_receipt_docx', kwargs={'pk': self.app1.pk})
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

