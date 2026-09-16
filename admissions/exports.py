import io
import urllib.parse
from django.http import HttpResponse, Http404
from django.utils import timezone
from django.shortcuts import get_object_or_404
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from admissions.models import EducationProgram, Application, ApplicationDocument, Specialty


def generate_competitive_list_workbook(program, financing_type=Application.FinancingType.BUDGET, only_originals=False, is_officer=False):
    """
    Генерирует объект Workbook openpyxl с детализированным конкурсным списком
    поступающих на заданную образовательную программу.
    
    Включает:
    - Официальную шапку МУ им. С.Ю. Витте с реквизитами программы и датой выгрузки;
    - Таблицу ранжированных абитуриентов с баллами по предметам и признаком оригинала;
    - Цветовую маркировку «зеленой зоны» (попадание в контрольные цифры приема);
    - Итоговую аналитическую сводку по конкурсу и текущему проходному баллу.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Конкурсный список"
    
    # Включаем сетку таблицы
    ws.views.sheetView[0].showGridLines = True

    specialty = program.specialty
    faculty = specialty.faculty
    places_count = specialty.budget_places if financing_type == Application.FinancingType.BUDGET else specialty.paid_places
    financing_display = "Бюджетная основа" if financing_type == Application.FinancingType.BUDGET else "Платная основа (договор)"

    # Выборка заявлений
    base_qs = Application.objects.filter(
        program=program,
        financing_type=financing_type,
    ).exclude(
        status__in=[Application.Status.DRAFT, Application.Status.REJECTED, Application.Status.WITHDRAWN]
    ).select_related(
        'applicant',
        'applicant__applicant_profile',
        'program__specialty',
        'program__specialty__faculty'
    ).prefetch_related(
        'exam_scores__subject',
        'documents'
    )

    candidates = []
    for app in base_qs:
        applicant = app.applicant
        profile = getattr(applicant, 'applicant_profile', None)

        if profile and profile.snils:
            snils_display = profile.snils.strip()
        else:
            snils_display = f"№ 2026-{app.id:04d}"

        scores = list(app.exam_scores.all())
        total_score = sum(s.score for s in scores)
        
        scores_str_list = [f"{s.subject.name}: {s.score}" for s in scores]
        scores_detail = ", ".join(scores_str_list) if scores_str_list else "Баллы на проверке"

        has_original = app.documents.filter(
            document_type__in=[ApplicationDocument.DocumentType.CERTIFICATE, ApplicationDocument.DocumentType.DIPLOMA]
        ).exists()

        fio = applicant.get_full_name() or applicant.username

        candidates.append({
            'app_id': app.id,
            'fio': fio,
            'snils': snils_display,
            'total_score': total_score,
            'scores_detail': scores_detail,
            'has_original': has_original,
            'status': app.get_status_display(),
            'submission_date': app.submission_date,
        })

    # Сортировка по убыванию баллов, затем по дате подачи
    candidates.sort(key=lambda x: (-x['total_score'], x['submission_date']))

    if only_originals:
        candidates = [c for c in candidates if c['has_original']]

    # Стили шрифтов и заливок
    font_title = Font(name='Calibri', size=14, bold=True, color='0F172A')
    font_subtitle = Font(name='Calibri', size=11, bold=True, color='1E3A8A')
    font_info = Font(name='Calibri', size=10, italic=False, color='334155')
    font_header = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
    font_data = Font(name='Calibri', size=10, color='1E293B')
    font_data_bold = Font(name='Calibri', size=10, bold=True, color='0F172A')
    font_summary_title = Font(name='Calibri', size=11, bold=True, color='1E3A8A')
    font_summary_label = Font(name='Calibri', size=10, bold=True, color='334155')
    font_summary_val = Font(name='Calibri', size=10, bold=True, color='0F172A')

    fill_header = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    fill_quota = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')      # Зеленая зона (в плане мест)
    fill_reserve = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')    # Резерв
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')      # Чередование строк
    fill_summary = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')    # Итоговый блок

    border_thin = Side(style='thin', color='CBD5E1')
    border_double = Side(style='double', color='1E3A8A')
    border_thick_bottom = Side(style='medium', color='1E3A8A')

    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    header_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thick_bottom)

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    # 1. ШАПКА ДОКУМЕНТА (Строки 1-5)
    ws['A1'] = "ЧОУ ВО «МОСКОВСКИЙ УНИВЕРСИТЕТ ИМЕНИ С.Ю. ВИТТЕ»"
    ws['A1'].font = font_title
    ws.merge_cells('A1:I1')
    ws['A1'].alignment = align_left

    ws['A2'] = "ОФИЦИАЛЬНЫЙ РАНЖИРОВАННЫЙ КОНКУРСНЫЙ СПИСОК ПОСТУПАЮЩИХ"
    ws['A2'].font = font_subtitle
    ws.merge_cells('A2:I2')
    ws['A2'].alignment = align_left

    ws['A3'] = f"Направление подготовки: {specialty.code} «{specialty.name}» ({specialty.get_education_level_display()})"
    ws['A3'].font = font_data_bold
    ws.merge_cells('A3:I3')
    ws['A3'].alignment = align_left

    ws['A4'] = f"Факультет: {faculty.name}  |  Форма обучения: {program.get_study_form_display()}  |  Основа: {financing_display}"
    ws['A4'].font = font_info
    ws.merge_cells('A4:I4')
    ws['A4'].alignment = align_left

    now_str = timezone.now().strftime('%d.%m.%Y в %H:%M')
    ws['A5'] = f"Контрольные цифры приема (план мест): {places_count}  |  Дата и время выгрузки: {now_str}"
    ws['A5'].font = font_info
    ws.merge_cells('A5:I5')
    ws['A5'].alignment = align_left

    # Строка 6 - разделитель
    ws.row_dimensions[6].height = 10

    # 2. ЗАГОЛОВКИ ТАБЛИЦЫ (Строка 7)
    headers = [
        ("№ п/п", 8, align_center),
        ("№ заявления", 14, align_center),
        ("СНИЛС / ID", 18, align_center),
        ("ФИО абитуриента", 32, align_left),
        ("Сумма баллов", 15, align_center),
        ("Предметы и баллы ЕГЭ / ВИ", 45, align_left),
        ("Оригинал аттестата", 18, align_center),
        ("Статус заявления", 24, align_center),
        ("Позиция в конкурсе", 22, align_center),
    ]

    header_row = 7
    ws.row_dimensions[header_row].height = 28

    for col_idx, (header_text, width, alignment) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = header_border
        
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 3. ДАННЫЕ АБИТУРИЕНТОВ (Строки 8+)
    current_row = 8
    total_candidates = len(candidates)
    originals_count = sum(1 for c in candidates if c['has_original'])

    for idx, candidate in enumerate(candidates, start=1):
        ws.row_dimensions[current_row].height = 22
        in_quota = (idx <= places_count) if places_count > 0 else False
        
        row_fill = fill_quota if in_quota else (fill_zebra if idx % 2 == 0 else fill_reserve)
        quota_label = "В зоне зачисления" if in_quota else "В резерве"
        original_label = "ДА (Оригинал)" if candidate['has_original'] else "НЕТ (Копия)"

        row_data = [
            (idx, align_center, font_data_bold if in_quota else font_data),
            (f"#{candidate['app_id']}", align_center, font_data),
            (candidate['snils'], align_center, font_data),
            (candidate['fio'], align_left, font_data_bold if in_quota else font_data),
            (candidate['total_score'], align_center, font_data_bold),
            (candidate['scores_detail'], align_left, font_data),
            (original_label, align_center, font_data_bold if candidate['has_original'] else font_data),
            (candidate['status'], align_center, font_data),
            (quota_label, align_center, font_data_bold if in_quota else font_data),
        ]

        for col_idx, (val, alignment, font_style) in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_style
            cell.fill = row_fill
            cell.alignment = alignment
            cell.border = cell_border

        current_row += 1

    # Если кандидатов нет
    if total_candidates == 0:
        ws.row_dimensions[current_row].height = 24
        cell = ws.cell(row=current_row, column=1, value="По выбранным критериям заявлений не найдено")
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
        cell.font = font_info
        cell.alignment = align_center
        cell.border = cell_border
        current_row += 1

    # 4. ИТОГОВАЯ СВОДКА ПО КОНКУРСУ (Строки N+2)
    current_row += 1  # пустая строка
    ws.row_dimensions[current_row].height = 12

    current_row += 1
    ws.row_dimensions[current_row].height = 24
    summary_title_cell = ws.cell(row=current_row, column=1, value="ИТОГОВЫЕ ПОКАЗАТЕЛИ КОНКУРСА:")
    summary_title_cell.font = font_summary_title
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)

    competition_ratio = round(total_candidates / places_count, 2) if places_count > 0 else 0.0
    if total_candidates > 0:
        if places_count > 0 and total_candidates >= places_count:
            passing_score = candidates[places_count - 1]['total_score']
        else:
            passing_score = candidates[-1]['total_score']
    else:
        passing_score = 0

    summary_items = [
        ("Всего заявлений в конкурсе:", f"{total_candidates} чел."),
        ("Подано оригиналов документов:", f"{originals_count} шт."),
        ("Контрольные цифры приема (мест):", f"{places_count} мест"),
        ("Конкурс на одно место:", f"{competition_ratio} чел./место"),
        ("Текущий расчетный проходной балл:", f"{passing_score} б."),
    ]

    for label, val in summary_items:
        current_row += 1
        ws.row_dimensions[current_row].height = 20
        
        c_lbl = ws.cell(row=current_row, column=1, value=label)
        c_lbl.font = font_summary_label
        c_lbl.alignment = align_left
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)

        c_val = ws.cell(row=current_row, column=4, value=val)
        c_val.font = font_summary_val
        c_val.alignment = align_left

    # Подписи ответственных лиц
    current_row += 2
    ws.row_dimensions[current_row].height = 22
    ws.cell(row=current_row, column=1, value="Ответственный секретарь приемной комиссии: ____________________ / ____________________ /").font = font_info
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)

    return wb


def export_rating_xlsx_response(program, financing_type=Application.FinancingType.BUDGET, only_originals=False, is_officer=False):
    """
    Генерирует Excel-файл и формирует HTTP-ответ с бинарным потоком .xlsx
    и правильными заголовками Content-Disposition для скачивания браузером.
    """
    wb = generate_competitive_list_workbook(
        program=program,
        financing_type=financing_type,
        only_originals=only_originals,
        is_officer=is_officer
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # Формируем имя файла
    spec_code_clean = program.specialty.code.replace('.', '_')
    basis_clean = "budget" if financing_type == Application.FinancingType.BUDGET else "paid"
    date_clean = timezone.now().strftime('%Y%m%d')
    filename = f"rating_{spec_code_clean}_{basis_clean}_{date_clean}.xlsx"

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    # Корректное кодирование для поддержки кириллицы и спецсимволов в имени файла
    encoded_filename = urllib.parse.quote(filename)
    response['Content-Disposition'] = f'attachment; filename="{filename}"; filename*=UTF-8\'\'{encoded_filename}'
    response['Cache-Control'] = 'no-cache'
    return response
